"""XLSX and PDF renderings of the sales report (same numbers as the web page)."""
from io import BytesIO

from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import services as S

NAVY, TEAL = '112C66', '009EA0'


def sections(f):
	"""[(title, rows)] shared by both formats."""
	return [
		(_("Bazeia rota"), S.by_route(f)), (_("Bazeia ró"), S.by_vessel(f)), (_("Bazeia operadór"), S.by_operator(f)),
		(_("Bazeia kategoria"), S.by_item(f)), (_("Bazeia pagamentu"), S.by_payment(f)),
		(_("Bazeia kaixa"), S.by_cashier(f)), (_("Loron-loron"), S.by_day(f)),
	]


def headers():
	return [_("Deskrisaun"), _("Bilhete"), _("Pasajeiru"), _("Bruto (USD)"), _("Fila osan (USD)"), _("Liquidu (USD)")]


def row_values(r):
	label = r['label'].strftime('%d/%m/%Y') if hasattr(r['label'], 'strftime') else r['label']
	return [label, r['tickets'], r['pax'], float(r['gross']), float(r['refunds']), float(r['net'])]


def summary_lines(f):
	s = S.summary(f)
	return [(_("Bilhete faan"), s['tickets']), (_("Pasajeiru"), s['pax']), (_("Bruto (USD)"), float(s['gross'])),
		(_("Fila osan (USD)"), float(s['refunds'])), (_("Liquidu (USD)"), float(s['net'])),
		(_("Bilhete kansela (sala) — la konta"), s['void_tickets'])]


def pdf_safe(text):
	"""The built-in PDF font has no arrow glyph (it would print as a black box)."""
	return str(text).replace('➔', '>')


def title(f):
	return f"{_('Relatóriu Vendas APORTIL')} {f.start:%d/%m/%Y} – {f.end:%d/%m/%Y}"


def to_xlsx(f):
	wb = Workbook()
	ws = wb.active
	ws.title = _("Rezumu")[:31]
	ws['A1'] = title(f)
	ws['A1'].font = Font(bold=True, size=14, color=NAVY)
	for i, (k, v) in enumerate(summary_lines(f), start=3):
		ws.cell(i, 1, k).font = Font(bold=True)
		ws.cell(i, 2, v)
	ws.column_dimensions['A'].width = 38
	ws.column_dimensions['B'].width = 16
	for name, rows in sections(f):
		sh = wb.create_sheet(name[:31])
		sh.append(headers())
		for c in sh[1]:
			c.font = Font(bold=True, color='FFFFFF')
			c.fill = PatternFill('solid', fgColor=NAVY)
			c.alignment = Alignment(horizontal='center')
		for r in rows:
			sh.append(row_values(r))
		for col, w in zip('ABCDEF', (44, 10, 11, 14, 16, 14)):
			sh.column_dimensions[col].width = w
		for row in sh.iter_rows(min_row=2, min_col=4, max_col=6):
			for c in row:
				c.number_format = '#,##0.00'
		sh.freeze_panes = 'A2'
	buf = BytesIO()
	wb.save(buf)
	return buf.getvalue()


def to_pdf(f):
	buf = BytesIO()
	doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
		title=title(f))
	st = getSampleStyleSheet()
	story = [Paragraph(f'<font color="#{NAVY}"><b>{title(f)}</b></font>', st['Title']), Spacer(1, 4 * mm)]
	t = Table([[k, f'{v:,.2f}' if isinstance(v, float) else str(v)] for k, v in summary_lines(f)], colWidths=[95 * mm, 40 * mm])
	t.setStyle(TableStyle([('FONTSIZE', (0, 0), (-1, -1), 10), ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
		('LINEBELOW', (0, 0), (-1, -1), 0.25, colors.lightgrey)]))
	story += [t, Spacer(1, 6 * mm)]
	for name, rows in sections(f):
		if not rows:
			continue
		story.append(Paragraph(f'<font color="#{TEAL}"><b>{name}</b></font>', st['Heading3']))
		data = [headers()] + [[(pdf_safe(c) if not isinstance(c, float) else f'{c:,.2f}') for c in row_values(r)] for r in rows]
		tb = Table(data, repeatRows=1, colWidths=[62 * mm, 16 * mm, 18 * mm, 24 * mm, 26 * mm, 24 * mm][:6])
		tb.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f'#{NAVY}')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
			('FONTSIZE', (0, 0), (-1, -1), 8), ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
			('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0f2f8')]), ('GRID', (0, 0), (-1, -1), 0.2, colors.lightgrey)]))
		story += [tb, Spacer(1, 4 * mm)]
	doc.build(story)
	return buf.getvalue()
