import io
from datetime import date, datetime, timedelta, timezone as dt_tz
from decimal import Decimal
from unittest import mock

from django.contrib.auth.models import Group
from django.core import mail
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from custom.models import FareItem, Operator, Port, Route, Vessel
from ticketing.models import ScanEvent, Sailing, SailingQuota, Ticket
from users.models import AuditLog, User

from . import emails, services as S
from .models import ReportLog

D = Decimal


def local(y, m, d, h=12, mi=0):
	return timezone.make_aware(datetime(y, m, d, h, mi))


class Base(TestCase):
	def setUp(self):
		call_command('setup_roles', verbosity=0)
		self.op1 = Operator.objects.create(name='Op1'); self.op2 = Operator.objects.create(name='Op2')
		self.v1 = Vessel.objects.create(operator=self.op1, name='V1'); self.v2 = Vessel.objects.create(operator=self.op2, name='V2')
		a, b, c = (Port.objects.create(code=x, name=n) for x, n in (('DIL', 'Dili'), ('BEL', 'Beloi'), ('MAH', 'Mahata')))
		self.r1 = Route.objects.create(origin=a, destination=b); self.r2 = Route.objects.create(origin=a, destination=c)
		self.adult = FareItem.objects.create(name_tet='Adultu', quota_group='passenger')
		self.buffalo = FareItem.objects.create(name_tet='Karau', quota_group='goods')
		self.s1 = self.sailing(self.v1, self.r1, local(2026, 10, 20)); self.s2 = self.sailing(self.v2, self.r2, local(2026, 10, 20, 15))
		self.admin = self.user('admin1', 'admin'); self.clerk = self.user('clerk', 'resepsionista'); self.clerk2 = self.user('clerk2', 'resepsionista')
		self.opuser = self.user('opuser', 'operator', operator=self.op1)

	def user(self, name, group, **kw):
		u = User.objects.create_user(name, password='x', **kw); u.groups.add(Group.objects.get(name=group)); return u

	def sailing(self, v, r, dep, **kw):
		s = Sailing.objects.create(vessel=v, route=r, departure=dep, **kw)
		for g, c in (('passenger', 100), ('vehicle', 5), ('goods', 50)):
			SailingQuota.objects.create(sailing=s, quota_group=g, capacity=c)
		return s

	def t(self, when, price='10', qty=1, sailing=None, item=None, status='issued', pay='cash', by=None, cancelled=None, **kw):
		x = Ticket.objects.create(sailing=sailing or self.s1, item=item or self.adult, unit_price=D(price), quantity=qty,
			status=status, payment_method=pay, sold_by=by or self.clerk, cancelled_at=cancelled, **kw)
		Ticket.objects.filter(pk=x.pk).update(created_at=when)
		return x

	def f(self, a, b, user=None, **kw):
		return S.Filters(user, a, b, **kw)


class MoneyRuleTests(Base):
	def test_gross_refund_void_net(self):
		self.t(local(2026, 10, 12), '10'); self.t(local(2026, 10, 12), '4', qty=3)           # 22 valid (+ pax 4)
		self.t(local(2026, 10, 12), '5', status='boarded')                                   # 5
		self.t(local(2026, 10, 12), '8', status='refunded', cancelled=local(2026, 10, 12, 18))  # refunded same day
		self.t(local(2026, 10, 12), '7', status='cancelled', cancelled=local(2026, 10, 12, 13))  # voided: not money
		s = S.summary(self.f(date(2026, 10, 12), date(2026, 10, 12)))
		self.assertEqual((s['gross'], s['refunds'], s['net'], s['voided']), (D('35'), D('8'), D('27'), D('7')))
		self.assertEqual((s['tickets'], s['pax'], s['refund_tickets'], s['void_tickets']), (4, 1 + 3 + 1 + 1, 1, 1))

	def test_goods_are_money_but_not_passengers(self):
		self.t(local(2026, 10, 12), '10', qty=2, item=self.buffalo)
		s = S.summary(self.f(date(2026, 10, 12), date(2026, 10, 12)))
		self.assertEqual((s['gross'], s['pax']), (D('20'), 0))

	def test_refund_is_dated_by_refund_not_sale(self):
		self.t(local(2026, 10, 12), '10', status='refunded', cancelled=local(2026, 10, 14))
		sale, refund = self.f(date(2026, 10, 12), date(2026, 10, 12)), self.f(date(2026, 10, 14), date(2026, 10, 14))
		self.assertEqual((S.summary(sale)['gross'], S.summary(sale)['net']), (D('10'), D('10')))   # day of sale: still counted
		self.assertEqual((S.summary(refund)['gross'], S.summary(refund)['net']), (D('0'), D('-10')))  # day of refund: money out
		whole = S.summary(self.f(date(2026, 10, 12), date(2026, 10, 14)))
		self.assertEqual(whole['net'], D('0'))

	def test_days_are_dili_days_not_utc(self):
		# 16:30 UTC on the 12th is 01:30 on the 13th in Dili (UTC+9)
		x = self.t(datetime(2026, 10, 12, 16, 30, tzinfo=dt_tz.utc), '10')
		self.assertEqual(S.summary(self.f(date(2026, 10, 12), date(2026, 10, 12)))['tickets'], 0)
		self.assertEqual(S.summary(self.f(date(2026, 10, 13), date(2026, 10, 13)))['tickets'], 1)
		self.assertEqual(S.by_day(self.f(date(2026, 10, 12), date(2026, 10, 14)))[0]['key'], date(2026, 10, 13))

	def test_range_is_inclusive_of_both_end_days(self):
		self.t(local(2026, 10, 12, 0, 0)); self.t(local(2026, 10, 14, 23, 59)); self.t(local(2026, 10, 15, 0, 0))
		self.assertEqual(S.summary(self.f(date(2026, 10, 12), date(2026, 10, 14)))['tickets'], 2)

	def test_breakdowns_add_up_and_include_refund_only_rows(self):
		self.t(local(2026, 10, 12), '10', sailing=self.s1, pay='cash')
		self.t(local(2026, 10, 12), '6', sailing=self.s2, pay='transfer', by=self.clerk2)
		self.t(local(2026, 10, 10), '9', sailing=self.s2, status='refunded', cancelled=local(2026, 10, 12), by=self.clerk2)  # sold before range
		f = self.f(date(2026, 10, 12), date(2026, 10, 12))
		total = S.summary(f)
		for rows in (S.by_route(f), S.by_vessel(f), S.by_operator(f), S.by_item(f), S.by_payment(f), S.by_cashier(f), S.by_day(f)):
			self.assertEqual(sum(r['net'] for r in rows), total['net'], rows)
			self.assertEqual(sum(r['gross'] for r in rows), total['gross'])
		r2 = next(r for r in S.by_route(f) if r['key'] == self.r2.pk)
		self.assertEqual((r2['gross'], r2['refunds'], r2['net'], r2['label']), (D('6'), D('9'), D('-3'), str(self.r2)))

	def test_daily_series_fills_gaps(self):
		self.t(local(2026, 10, 12), '10'); self.t(local(2026, 10, 14), '4')
		ser = S.daily_series(self.f(date(2026, 10, 12), date(2026, 10, 15)))
		self.assertEqual([(r['day'].day, r['net'], r['tickets']) for r in ser], [(12, D('10'), 1), (13, D('0'), 0), (14, D('4'), 1), (15, D('0'), 0)])

	def test_period_bounds(self):
		self.assertEqual(S.period_bounds('daily', date(2026, 10, 12)), (date(2026, 10, 12),) * 2)
		self.assertEqual(S.period_bounds('monthly', date(2026, 2, 10)), (date(2026, 2, 1), date(2026, 2, 28)))
		self.assertEqual(S.period_bounds('monthly', date(2028, 2, 10))[1], date(2028, 2, 29))
		self.assertEqual(S.period_bounds('yearly', date(2026, 6, 1)), (date(2026, 1, 1), date(2026, 12, 31)))


class ScopeAndAccessTests(Base):
	def setUp(self):
		super().setUp()
		self.t(local(2026, 10, 12), '10', sailing=self.s1); self.t(local(2026, 10, 12), '30', sailing=self.s2)
		self.rng = {'start': '2026-10-01', 'end': '2026-10-31'}

	def test_operator_sees_only_own_money_everywhere(self):
		f = self.f(date(2026, 10, 1), date(2026, 10, 31), user=self.opuser)
		self.assertEqual(S.summary(f)['gross'], D('10'))
		self.assertEqual([r['label'] for r in S.by_vessel(f)], ['V1 (Op1)'])
		self.client.force_login(self.opuser)
		r = self.client.get(reverse('report-sales'), self.rng)
		self.assertEqual(r.context['summary']['gross'], D('10'))
		self.assertNotIn('operator', r.context['form'].fields)
		# trying to peek at another operator through the query string changes nothing
		r = self.client.get(reverse('report-sales'), {**self.rng, 'operator': self.op2.pk})
		self.assertEqual(r.context['summary']['gross'], D('10'))
		api = self.client.get('/api/report/venda/rota/', {'from': '2026-10-01', 'to': '2026-10-31'}).json()
		self.assertEqual(api['label'], [str(self.r1)])

	def test_operator_cannot_filter_to_foreign_vessel(self):
		self.client.force_login(self.opuser)
		r = self.client.get(reverse('report-sales'), {**self.rng, 'vessel': self.v2.pk})
		self.assertFalse(r.context['form'].is_valid())  # not in the allowed choices

	def test_admin_sees_everything_and_can_filter_operator(self):
		self.client.force_login(self.admin)
		self.assertEqual(self.client.get(reverse('report-sales'), self.rng).context['summary']['gross'], D('40'))
		self.assertEqual(self.client.get(reverse('report-sales'), {**self.rng, 'operator': self.op2.pk}).context['summary']['gross'], D('30'))

	def test_permissions(self):
		pages = [reverse('report-dash'), reverse('report-sales'), reverse('report-boarding'), '/api/report/venda/rota/']
		self.client.force_login(self.clerk)
		for p in pages:
			self.assertEqual(self.client.get(p).status_code, 403, p)
		self.assertEqual(self.client.get(reverse('report-cashier')).status_code, 200)
		self.client.logout()
		for p in pages[:3]:
			self.assertEqual(self.client.get(p).status_code, 302, p)
		self.assertEqual(self.client.get('/api/report/venda/rota/').status_code, 403)
		self.client.force_login(self.opuser)
		self.assertEqual(self.client.get(reverse('report-cashier')).status_code, 403)  # counter cash is not an operator's business
		for p in pages:
			self.assertEqual(self.client.get(p).status_code, 200, p)

	def test_filter_validation_blocks_page_and_export(self):
		self.client.force_login(self.admin)
		r = self.client.get(reverse('report-sales'), {'start': '2026-10-31', 'end': '2026-10-01'})
		self.assertFalse(r.context['ok']); self.assertIsNone(r.context['summary'])
		r = self.client.get(reverse('report-sales'), {'start': '2026-10-31', 'end': '2026-10-01', 'fmt': 'pdf'})
		self.assertEqual(r.status_code, 200); self.assertNotEqual(r['Content-Type'], 'application/pdf')
		r = self.client.get(reverse('report-sales'), {'start': '2020-01-01', 'end': '2026-10-01'})
		self.assertFalse(r.context['ok'])

	def test_api_defaults_params_and_bad_input(self):
		self.client.force_login(self.admin)
		d = self.client.get('/api/report/venda/diaria/').json()
		self.assertEqual((len(d['label']), set(d)), (30, {'label', 'obj', 'obj2'}))  # 30 zero-filled days
		d = self.client.get('/api/report/venda/diaria/', {'from': '2026-10-10', 'to': '2026-10-13'}).json()
		self.assertEqual((d['label'], d['obj'], d['obj2']), (['10/10', '11/10', '12/10', '13/10'], [0, 0, 40.0, 0], [0, 0, 2, 0]))
		for bad in ({'from': 'x', 'to': 'y'}, {'from': '2026-10-13', 'to': '2026-10-01'}, {'from': '2000-01-01', 'to': '2026-10-01'}):
			r = self.client.get('/api/report/venda/diaria/', bad)
			self.assertEqual((r.status_code, len(r.json()['label'])), (200, 30), bad)
		for name in ('rota', 'navio', 'pagamentu', 'kategoria'):
			j = self.client.get(f'/api/report/venda/{name}/', {'from': '2026-10-01', 'to': '2026-10-31'}).json()
			self.assertEqual(sum(j['obj']), 40.0, name)


class CashierTests(Base):
	def test_drawer_math_and_scope(self):
		self.t(local(2026, 10, 12), '10', pay='cash'); self.t(local(2026, 10, 12), '20', pay='transfer')
		self.t(local(2026, 10, 12), '6', pay='cash', status='refunded', cancelled=local(2026, 10, 12, 17))
		self.t(local(2026, 10, 12), '50', pay='cash', by=self.clerk2)
		rows = S.cashier_day(self.admin, date(2026, 10, 12))
		mine = next(r for r in rows if r['user'] == self.clerk)
		self.assertEqual((mine['cash'], mine['transfer'], mine['refunds'], mine['drawer'], mine['tickets']), (D('16'), D('20'), D('6'), D('10'), 3))
		self.assertEqual(len(S.cashier_day(self.clerk, date(2026, 10, 12), cashier=self.clerk)), 1)
		self.client.force_login(self.clerk)
		r = self.client.get(reverse('report-cashier'), {'date': '2026-10-12'})
		self.assertEqual([x['user'] for x in r.context['rows']], [self.clerk]); self.assertTrue(r.context['own_only'])
		self.assertNotContains(r, 'clerk2')
		self.client.force_login(self.admin)
		r = self.client.get(reverse('report-cashier'), {'date': '2026-10-12'})
		self.assertEqual(len(r.context['rows']), 2); self.assertEqual(r.context['totals']['drawer'], D('60'))
		self.assertEqual(self.client.get(reverse('report-cashier'), {'date': 'garbage'}).status_code, 200)

	def test_cash_refund_of_earlier_day_reduces_that_days_drawer(self):
		self.t(local(2026, 10, 10), '9', pay='cash', status='refunded', cancelled=local(2026, 10, 12, 9))
		rows = S.cashier_day(self.admin, date(2026, 10, 12))
		self.assertEqual((rows[0]['cash'], rows[0]['drawer']), (D('0'), D('-9')))


class BoardingTests(Base):
	def test_counts_no_shows_only_after_departure(self):
		for _ in range(3): self.t(local(2026, 10, 12), sailing=self.s1)
		Ticket.objects.filter(pk=Ticket.objects.filter(sailing=self.s1).first().pk).update(status='boarded')
		self.t(local(2026, 10, 12), sailing=self.s1, status='cancelled')
		ScanEvent.objects.create(sailing=self.s1, serial='x', result='duplicate', scanned_at=timezone.now())
		ScanEvent.objects.create(sailing=self.s1, serial='y', result='invalid', scanned_at=timezone.now())
		ScanEvent.objects.create(sailing=self.s1, serial='z', result='wrong_sailing', scanned_at=timezone.now())
		rows = S.boarding(self.admin, date(2026, 10, 20), date(2026, 10, 20))
		r = next(r for r in rows if r['sailing'] == self.s1)
		self.assertEqual((r['tickets'], r['boarded'], r['no_show'], r['duplicates'], r['invalid'], r['wrong']), (3, 1, None, 1, 1, 1))
		Sailing.objects.filter(pk=self.s1.pk).update(status='departed')
		r = next(r for r in S.boarding(self.admin, date(2026, 10, 20), date(2026, 10, 20)) if r['sailing'].pk == self.s1.pk)
		self.assertEqual(r['no_show'], 2)
		self.assertEqual(len(S.boarding(self.opuser, date(2026, 10, 20), date(2026, 10, 20))), 1)  # operator: own sailings only
		self.client.force_login(self.admin)
		self.assertEqual(self.client.get(reverse('report-boarding'), {'start': '2026-10-20', 'end': '2026-10-20'}).status_code, 200)


class ExportTests(Base):
	def setUp(self):
		super().setUp()
		self.t(local(2026, 10, 12), '10', sailing=self.s1); self.t(local(2026, 10, 12), '30', qty=2, sailing=self.s2, pay='transfer')
		self.client.force_login(self.admin)
		self.q = {'start': '2026-10-01', 'end': '2026-10-31'}

	def test_xlsx_has_sheets_and_right_numbers(self):
		r = self.client.get(reverse('report-sales'), {**self.q, 'fmt': 'xlsx'})
		self.assertEqual(r.status_code, 200)
		self.assertIn('spreadsheetml', r['Content-Type']); self.assertIn('aportil-vendas-20261001-20261031.xlsx', r['Content-Disposition'])
		wb = load_workbook(io.BytesIO(r.content))
		self.assertEqual(len(wb.sheetnames), 8)
		summary = {row[0].value: row[1].value for row in wb.worksheets[0].iter_rows(min_row=3, max_col=2) if row[0].value}
		self.assertEqual(summary['Liquidu (USD)'], 70.0)
		route_sheet = wb.worksheets[1]
		self.assertEqual(sorted(float(c.value) for c in route_sheet['F'][1:]), [10.0, 60.0])
		self.assertTrue(AuditLog.objects.filter(action='report_export').exists())

	def test_pdf_is_a_pdf_and_follows_language(self):
		r = self.client.get(reverse('report-sales'), {**self.q, 'fmt': 'pdf'})
		self.assertEqual((r['Content-Type'], r.content[:5]), ('application/pdf', b'%PDF-'))
		self.assertGreater(len(r.content), 1500)

	def test_pdf_labels_avoid_glyphs_the_font_lacks(self):
		from . import exports
		self.assertEqual(exports.pdf_safe('Dili ➔ Beloi'), 'Dili > Beloi')

	def test_operator_export_contains_only_own_data(self):
		self.client.force_login(self.opuser)
		wb = load_workbook(io.BytesIO(self.client.get(reverse('report-sales'), {**self.q, 'fmt': 'xlsx'}).content))
		self.assertEqual(sorted(float(c.value) for c in wb.worksheets[1]['F'][1:]), [10.0])

	def test_clerk_cannot_export(self):
		self.client.force_login(self.clerk)
		self.assertEqual(self.client.get(reverse('report-sales'), {**self.q, 'fmt': 'xlsx'}).status_code, 403)


@override_settings(REPORT_RECIPIENTS=['boss@example.com', 'ops@example.com'])
class EmailTests(Base):
	def setUp(self):
		super().setUp()
		self.t(local(2026, 10, 12), '10', sailing=self.s1); self.t(local(2026, 10, 12), '30', qty=2, sailing=self.s2)
		self.t(local(2026, 10, 12), '4', status='refunded', cancelled=local(2026, 10, 12, 20))

	def test_build_is_portuguese_and_correct(self):
		key, subject, body, files = emails.build('daily', date(2026, 10, 12))
		self.assertEqual(key, '2026-10-12')
		self.assertIn('Relatório de vendas diário — 2026-10-12', subject)
		self.assertIn('Bilhetes vendidos: 3', body); self.assertIn('Líquido (USD): 70.00', body); self.assertIn('Reembolsos (USD): 4.00', body)
		self.assertEqual([n for n, *_ in files], ['aportil-vendas-2026-10-12.pdf', 'aportil-vendas-2026-10-12.xlsx'])

	def test_send_attaches_files_and_goes_to_manager_not_operators(self):
		log = emails.send('daily', date(2026, 10, 12))
		self.assertTrue(log.ok)
		m = mail.outbox[0]
		self.assertEqual(m.to, ['boss@example.com', 'ops@example.com'])
		self.assertEqual(sorted(a[0] for a in m.attachments), ['aportil-vendas-2026-10-12.pdf', 'aportil-vendas-2026-10-12.xlsx'])
		self.assertTrue(next(a for a in m.attachments if a[0].endswith('.pdf'))[1][:4] == b'%PDF')

	def test_idempotent_force_and_failed_attempt_can_be_retried(self):
		emails.send('daily', date(2026, 10, 12))
		self.assertIsNone(emails.send('daily', date(2026, 10, 12))); self.assertEqual(len(mail.outbox), 1)
		emails.send('daily', date(2026, 10, 12), force=True); self.assertEqual(len(mail.outbox), 2)
		with mock.patch('django.core.mail.EmailMessage.send', side_effect=OSError('smtp down')):
			bad = emails.send('daily', date(2026, 10, 13))
		self.assertFalse(bad.ok); self.assertIn('smtp down', bad.error)
		self.assertTrue(emails.send('daily', date(2026, 10, 13)).ok)  # a failure never blocks the retry

	@override_settings(REPORT_RECIPIENTS=[])
	def test_no_recipients_is_logged_not_sent(self):
		log = emails.send('daily', date(2026, 10, 12))
		self.assertFalse(log.ok); self.assertIn('no recipients', log.error); self.assertEqual(len(mail.outbox), 0)

	def test_command_defaults_and_errors(self):
		out = io.StringIO()
		with mock.patch('django.utils.timezone.localdate', return_value=date(2026, 11, 3)):
			call_command('send_reports', period='monthly', stdout=out)   # previous month = October
			call_command('send_reports', period='yearly', stdout=out)    # previous year = 2025
		self.assertEqual(set(ReportLog.objects.values_list('period', 'key')), {('monthly', '2026-10'), ('yearly', '2025')})
		self.assertIn('Líquido (USD): 70.00', mail.outbox[0].body)
		call_command('send_reports', period='daily', date='2026-10-12', stdout=out)
		call_command('send_reports', period='daily', date='2026-10-12', stdout=out)
		self.assertIn('Already sent', out.getvalue())
		with self.assertRaises(CommandError):
			call_command('send_reports', period='daily', date='nope')
		with mock.patch('django.core.mail.EmailMessage.send', side_effect=OSError('x')), self.assertRaises(CommandError):
			call_command('send_reports', period='daily', date='2026-10-11')
